from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Iterable, Sequence

from aiohttp import web
from server import PromptServer


routes = PromptServer.instance.routes

MAX_DOCUMENT_SIZE = 20 * 1024 * 1024
MAX_ARCHIVE_SIZE = 128 * 1024 * 1024
MAX_ROWS = 500
MAX_COLUMNS = 64
SUPPORTED_EXTENSIONS = {".docx", ".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".md", ".pdf"}


class StoryboardImportError(ValueError):
    pass


_HEADER_ALIASES = {
    "shot_no": {
        "镜号", "镜头号", "分镜号", "序号", "编号", "shot", "shotno",
        "shotnumber", "scene", "sceneno",
    },
    "time_range": {
        "时长", "时间", "时间段", "时间码", "起止时间", "time", "timerange",
        "timecode", "duration",
    },
    "image_prompt": {
        "画面内容", "画面描述", "画面", "提示词", "图像提示词", "分镜内容",
        "镜头内容", "视觉描述", "visual", "visualdescription", "imageprompt",
        "prompt", "description",
    },
    "camera_notes": {
        "镜头备注", "镜头说明", "运镜", "镜头运动", "拍摄备注", "摄影备注",
        "备注", "cameranotes", "cameramovement", "notes", "note",
    },
    "reference_image": {
        "参考图", "参考图片", "参考图像", "参考素材", "referenceimage",
        "refimage", "reference", "image",
    },
}

_IGNORED_HEADER_ALIASES = {
    "配音旁白", "配音", "旁白", "对白", "台词", "文案", "字幕",
    "voiceover", "narration", "dialogue", "dialog", "audio",
}


def _clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(part.strip() for part in text.split("\n")).strip()


def _header_key(value: object) -> str:
    text = unicodedata.normalize("NFKC", _clean(value)).casefold()
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", text)


def _header_map(row: Sequence[object]) -> dict[str, int]:
    matched: dict[str, int] = {}
    for index, cell in enumerate(row[:MAX_COLUMNS]):
        key = _header_key(cell)
        if not key or key in _IGNORED_HEADER_ALIASES:
            continue
        for field, aliases in _HEADER_ALIASES.items():
            if field not in matched and key in aliases:
                matched[field] = index
                break
    return matched


def _seconds(token: str) -> float | None:
    token = token.strip()
    if not token:
        return None
    try:
        if ":" not in token:
            return float(token)
        parts = [float(part) for part in token.split(":")]
    except ValueError:
        return None
    total = 0.0
    for part in parts:
        total = total * 60 + part
    return total


def duration_from_time_range(value: str) -> int:
    text = unicodedata.normalize("NFKC", value).strip().lower()
    if not text:
        return 3
    token = r"\d+(?::\d{1,2}){0,2}(?:\.\d+)?"
    range_match = re.search(rf"({token})\s*(?:-|~|至|到)\s*({token})", text)
    if range_match:
        start = _seconds(range_match.group(1))
        end = _seconds(range_match.group(2))
        if start is not None and end is not None and end > start:
            return max(1, min(3600, round(end - start)))
    single = re.search(rf"({token})\s*(?:s|秒|sec|seconds?)?\s*$", text)
    if single:
        seconds = _seconds(single.group(1))
        if seconds is not None and seconds > 0:
            return max(1, min(3600, round(seconds)))
    return 3


def _archive_preflight(payload: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            total = sum(info.file_size for info in archive.infolist())
            if total > MAX_ARCHIVE_SIZE:
                raise StoryboardImportError("document expands beyond the 128 MB safety limit")
    except zipfile.BadZipFile as exc:
        raise StoryboardImportError("document is not a valid Office archive") from exc


def _trim_table(rows: Iterable[Sequence[object]]) -> list[list[str]]:
    result: list[list[str]] = []
    for source_index, row in enumerate(rows):
        if source_index >= MAX_ROWS:
            break
        cells = [_clean(cell) for cell in list(row)[:MAX_COLUMNS]]
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            result.append(cells)
    return result


def _best_table(tables: Sequence[list[list[str]]]):
    candidates = []
    for table in tables:
        for header_index, row in enumerate(table[:12]):
            mapping = _header_map(row)
            if "image_prompt" in mapping:
                candidates.append((5 + len(mapping), len(table) - header_index, table, header_index, mapping))
    if not candidates:
        return None
    _, _, table, header_index, mapping = max(candidates, key=lambda item: (item[0], item[1]))
    return table, header_index, mapping


def _cell(row: Sequence[str], mapping: dict[str, int], field: str) -> str:
    index = mapping.get(field)
    return row[index] if index is not None and index < len(row) else ""


def _table_to_shots(table, header_index, mapping) -> tuple[list[dict], list[str]]:
    shots: list[dict] = []
    warnings: list[str] = []
    for source_row, row in enumerate(table[header_index + 1 :], start=header_index + 2):
        non_empty = [cell for cell in row if cell]
        if non_empty and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in non_empty):
            continue
        prompt = _cell(row, mapping, "image_prompt").strip()
        if not any(cell.strip() for cell in row):
            continue
        if not prompt:
            warnings.append(f"Skipped row {source_row}: no visual prompt")
            continue
        time_range = _cell(row, mapping, "time_range")
        reference = _cell(row, mapping, "reference_image")
        image_url = reference if reference.startswith("/view?") else ""
        if reference and not image_url:
            warnings.append(
                f"Row {source_row}: reference image path was not imported; upload it in ComfyUI"
            )
        shots.append(
            {
                "shot_no": _cell(row, mapping, "shot_no") or str(len(shots) + 1),
                "time_range": time_range,
                "duration": duration_from_time_range(time_range),
                "prompt": prompt,
                "image_prompt": prompt,
                "camera_notes": _cell(row, mapping, "camera_notes"),
                "image_url": image_url,
            }
        )
    if not shots:
        raise StoryboardImportError("the detected table has no rows with visual prompts")
    return shots, warnings


def _numbered_list_to_table(lines: Iterable[str]):
    rows = [["镜号", "画面内容"]]
    for line in lines:
        match = re.match(r"^\s*(?:(\d+)\s*[.)、]|[-*•])\s*(.+?)\s*$", line)
        if match:
            rows.append([match.group(1) or str(len(rows)), match.group(2)])
    return rows if len(rows) > 1 else None


def _decode_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise StoryboardImportError("text encoding is not supported")


def _text_tables(payload: bytes, suffix: str):
    text = _decode_text(payload)
    lines = text.splitlines()
    tables = []
    if suffix == ".md" or any("|" in line for line in lines[:20]):
        pipe_rows = [
            [cell.strip() for cell in line.strip().strip("|").split("|")]
            for line in lines
            if "|" in line
        ]
        trimmed = _trim_table(pipe_rows)
        if trimmed:
            tables.append(trimmed)
    delimiter = "\t" if suffix == ".tsv" else ","
    if suffix not in {".tsv", ".csv"} and any("\t" in line for line in lines[:20]):
        delimiter = "\t"
    if suffix in {".csv", ".tsv"} or delimiter == "\t":
        parsed = _trim_table(csv.reader(io.StringIO(text), delimiter=delimiter))
        if parsed:
            tables.append(parsed)
    numbered = _numbered_list_to_table(lines)
    if numbered:
        tables.append(numbered)
    title = next((line.strip().lstrip("#").strip() for line in lines if line.strip()), "")
    return tables, title


def _docx_tables(payload: bytes):
    _archive_preflight(payload)
    try:
        from docx import Document

        document = Document(io.BytesIO(payload))
    except (ImportError, ValueError, KeyError) as exc:
        raise StoryboardImportError(f"could not read DOCX: {exc}") from exc
    tables = [
        _trim_table([[cell.text for cell in row.cells] for row in table.rows])
        for table in document.tables
    ]
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    numbered = _numbered_list_to_table(paragraphs)
    if numbered:
        tables.append(numbered)
    return [table for table in tables if table], (paragraphs[0] if paragraphs else "")


def _xlsx_tables(payload: bytes):
    _archive_preflight(payload)
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
    except (ImportError, ValueError, KeyError, OSError) as exc:
        raise StoryboardImportError(f"could not read workbook: {exc}") from exc
    try:
        tables = [_trim_table(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets]
        title = workbook.worksheets[0].title if workbook.worksheets else ""
        return [table for table in tables if table], title
    finally:
        workbook.close()


def _pdf_tables(payload: bytes):
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(payload)) as pdf:
            tables = []
            lines = []
            for page in pdf.pages[:100]:
                for table in page.extract_tables() or []:
                    trimmed = _trim_table(table)
                    if trimmed:
                        tables.append(trimmed)
                lines.extend((page.extract_text() or "").splitlines())
    except (ImportError, ValueError, OSError) as exc:
        raise StoryboardImportError(f"could not read PDF: {exc}") from exc
    numbered = _numbered_list_to_table(lines)
    if numbered:
        tables.append(numbered)
    return tables, next((line.strip() for line in lines if line.strip()), "")


def parse_storyboard_document(filename: str, payload: bytes) -> dict:
    safe_name = Path(filename).name
    suffix = Path(safe_name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise StoryboardImportError(f"unsupported document type {suffix!r}; use {allowed}")
    if not payload:
        raise StoryboardImportError("document is empty")
    if len(payload) > MAX_DOCUMENT_SIZE:
        raise StoryboardImportError("document is larger than 20 MB")

    if suffix == ".docx":
        tables, title = _docx_tables(payload)
    elif suffix in {".xlsx", ".xlsm"}:
        tables, title = _xlsx_tables(payload)
    elif suffix == ".pdf":
        tables, title = _pdf_tables(payload)
    else:
        tables, title = _text_tables(payload, suffix)

    best = _best_table(tables)
    if best is None:
        raise StoryboardImportError(
            "no storyboard table found; include a visual column such as 画面内容 or 提示词"
        )
    table, header_index, mapping = best
    shots, warnings = _table_to_shots(table, header_index, mapping)
    return {
        "filename": safe_name,
        "document_title": title,
        "detected_columns": list(mapping),
        "shots": shots,
        "warnings": warnings[:50],
    }


@routes.post("/daelab/storyboard/import_document")
async def import_storyboard_document(request: web.Request) -> web.Response:
    try:
        reader = await request.multipart()
    except (AssertionError, ValueError) as exc:
        return web.json_response({"error": f"expected multipart body: {exc}"}, status=400)

    field = None
    while True:
        part = await reader.next()
        if part is None:
            break
        if part.name == "file":
            field = part
            break
    if field is None:
        return web.json_response({"error": "expected multipart field 'file'"}, status=400)

    payload = bytearray()
    while True:
        chunk = await field.read_chunk()
        if not chunk:
            break
        payload.extend(chunk)
        if len(payload) > MAX_DOCUMENT_SIZE:
            return web.json_response({"error": "document is larger than 20 MB"}, status=400)

    try:
        result = parse_storyboard_document(field.filename or "", bytes(payload))
    except StoryboardImportError as exc:
        return web.json_response({"error": str(exc)}, status=400)
    except Exception:
        return web.json_response({"error": "document could not be parsed"}, status=400)
    return web.json_response(result)
