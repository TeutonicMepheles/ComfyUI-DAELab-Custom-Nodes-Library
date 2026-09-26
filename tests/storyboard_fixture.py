"""Synthetic multi-table Word fixture for local browser acceptance."""
import io
import sys
from docx import Document
from PIL import Image

doc = Document()
doc.add_paragraph("分镜导入验收样例（合成素材）")
for title, count in [("基础表", 1), ("分镜表", 3)]:
    doc.add_paragraph(title)
    table = doc.add_table(rows=1, cols=6)
    for cell, label in zip(table.rows[0].cells, ["镜号", "时长", "画面内容", "旁白", "运镜", "参考图"]):
        cell.text = label
    for i in range(count):
        cells = table.add_row().cells
        for cell, value in zip(cells, [str(i + 1), "待确认" if i == 1 else "5s", "" if i == 2 else "星空中的飞船", "这是旁白原文", "缓慢推进", ""]):
            cell.text = value
        image = io.BytesIO()
        Image.new("RGB", (120, 80), ["#4689aa", "#ddab54", "#ab6789"][i]).save(image, format="PNG")
        image.seek(0)
        cells[5].paragraphs[0].add_run().add_picture(image)
doc.save(sys.argv[1])
