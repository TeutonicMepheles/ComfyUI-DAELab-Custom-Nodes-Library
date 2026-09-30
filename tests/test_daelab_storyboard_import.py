from __future__ import annotations

import importlib.util
import pathlib
import sys
import types
import unittest
from unittest import mock


class RoutesStub:
    def post(self, _path):
        return lambda function: function


server_stub = types.ModuleType("server")
server_stub.PromptServer = types.SimpleNamespace(
    instance=types.SimpleNamespace(routes=RoutesStub())
)

path = pathlib.Path(__file__).parents[1] / "nodes" / "daelab_comfytv_storyboard" / "document_import.py"
spec = importlib.util.spec_from_file_location("daelab_storyboard_import", path)
module = importlib.util.module_from_spec(spec)
with mock.patch.dict(sys.modules, {"server": server_stub}):
    spec.loader.exec_module(module)


class DAELabStoryboardImportTests(unittest.TestCase):
    def test_explicit_description_beats_picture_column_and_decimal_is_not_list(self):
        self.assertEqual(module._header_map(['序号','画面','描述','备注'])['image_prompt'], 2)
        self.assertEqual(module._header_map(['段落',' 设计描述','画面参考','备注'])['image_prompt'], 1)
        self.assertIsNone(module._numbered_list_to_table(['3.8米']))
        self.assertEqual(module._numbered_list_to_table(['3. 开始发射'])[1], ['3','开始发射'])

    def test_docx_preview_preserves_tables_rows_images_and_original_text(self):
        import io
        from docx import Document
        from PIL import Image
        doc = Document()
        for _ in range(2):
            table = doc.add_table(rows=1, cols=4)
            for cell, value in zip(table.rows[0].cells, ["镜号", "画面内容", "旁白", "参考图"]):
                cell.text = value
            cells = table.add_row().cells
            cells[0].text = "01"
            cells[2].text = "保留旁白"
            blob = io.BytesIO()
            Image.new("RGB", (4, 4), "red").save(blob, format="PNG")
            blob.seek(0)
            cells[3].paragraphs[0].add_run().add_picture(blob)
        payload = io.BytesIO()
        doc.save(payload)
        result = module.preview_storyboard_document("test.docx", payload.getvalue())
        self.assertEqual(len(result["tables"]), 2)
        self.assertEqual(result["tables"][0]["rows"][1][2], "保留旁白")
        self.assertEqual(len(result["assets"]), 1)
        self.assertEqual(result["tables"][1]["images"]["1:3"], [result["assets"][0]["id"]])
        self.assertNotIn("dialogue", result["tables"][0]["mapping"])

    def test_preview_allows_unknown_headers_for_manual_mapping(self):
        result = module.preview_storyboard_document("test.csv", b"a,b\n1,hello\n")
        self.assertEqual(result["tables"][0]["rows"][1], ["1", "hello"])

    def test_docx_tracked_insertions_and_deletions_use_final_view(self):
        import io
        from docx import Document
        from docx.oxml import OxmlElement
        from PIL import Image
        doc = Document()
        para = doc.add_paragraph('最终说明')
        table = doc.add_table(rows=3, cols=3)
        for cell, text in zip(table.rows[0].cells, ['镜号', '画面内容', '旁白']):
            cell.text = text
        for cell, text in zip(table.rows[1].cells, ['01', '最终画面', '旁白另存']):
            cell.text = text
        table.rows[2].cells[1].text = '删除整行'
        table.rows[2]._tr.get_or_add_trPr().append(OxmlElement('w:del'))
        blob = io.BytesIO()
        Image.new('RGB', (3, 3), 'blue').save(blob, format='PNG')
        blob.seek(0)
        table.rows[1].cells[1].paragraphs[0].add_run().add_picture(blob)
        # Both paragraph- and run-level insertion wrappers occur in real Word files.
        for p in [p for row in table.rows for cell in row.cells for p in cell.paragraphs]:
            insertion = OxmlElement('w:ins')
            for run in list(p._p.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r')):
                insertion.append(run)
            p._p.append(insertion)
        insertion = OxmlElement('w:ins')
        para._p.addprevious(insertion)
        insertion.append(para._p)
        deleted = OxmlElement('w:del')
        run = OxmlElement('w:r')
        old = OxmlElement('w:delText'); old.text = '旧画面不可混入'
        run.append(old); deleted.append(run)
        table.rows[1].cells[1].paragraphs[0]._p.append(deleted)
        payload = io.BytesIO(); doc.save(payload)
        result = module.preview_storyboard_document('revision.docx', payload.getvalue())
        self.assertEqual(result['tables'][0]['rows'], [['镜号','画面内容','旁白'],['01','最终画面','旁白另存']])
        self.assertEqual(result['document_notes'], '最终说明')
        self.assertEqual(len(result['tables'][0]['images']['1:1']), 1)
        self.assertTrue(any('修订' in warning for warning in result['warnings']))
        self.assertEqual(module.parse_storyboard_document('revision.docx', payload.getvalue())['shots'][0]['image_prompt'], '最终画面')

    def test_business_csv_ignores_voiceover_and_allows_blank_notes(self):
        payload = (
            "镜号,时长,画面内容,配音旁白,镜头备注,参考图\n"
            "01,00-08s,浩瀚星空铺开,每一次火箭冲破云层,,/view?filename=space.png&type=input\n"
            "02,08-18s,镜头下移,另一段旁白,纵深运镜,\n"
        ).encode("utf-8")
        result = module.parse_storyboard_document("分镜表.csv", payload)
        self.assertEqual(len(result["shots"]), 2)
        self.assertNotIn("dialogue", result["shots"][0])
        self.assertEqual(result["shots"][0]["camera_notes"], "")
        self.assertEqual(result["shots"][0]["image_url"], "/view?filename=space.png&type=input")
        self.assertEqual(result["shots"][1]["duration"], 10)

    def test_requires_visual_prompt_column(self):
        with self.assertRaisesRegex(module.StoryboardImportError, "no storyboard table"):
            module.parse_storyboard_document("bad.csv", "镜号,旁白\n1,你好\n".encode("utf-8"))


if __name__ == "__main__":
    unittest.main()
